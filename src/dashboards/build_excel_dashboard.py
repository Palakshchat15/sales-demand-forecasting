"""
Builds outputs/Sales_Forecast_Dashboard.xlsx from the CSVs in data/processed/
(written by forecast_prophet.py and export_dashboard_data.py).

Sheets: Dashboard (KPI cards + 4 charts), Summary (the aggregates the charts
plot), and one data sheet per source table. KPI cards are Excel formulas over
the data sheets, so they recalculate if the data sheets are refreshed.
"""
import os
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

PROJECT_ROOT = Path(__file__).resolve().parents[2]
PROCESSED_DIR = Path(os.environ.get("DASHBOARD_DATA_DIR", PROJECT_ROOT / "data" / "processed"))
OUTPUT_PATH = PROJECT_ROOT / "outputs" / "Sales_Forecast_Dashboard.xlsx"

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
KPI_FILL = PatternFill(start_color="EFF6FC", end_color="EFF6FC", fill_type="solid")
KPI_LABEL_FONT = Font(size=10, color="44546A", bold=True)
KPI_VALUE_FONT = Font(size=18, bold=True, color="1F4E78")


def write_table(ws, df, name, start_row=1, start_col=1):
    """Writes df as a formatted Excel Table; returns (first_row, last_row)."""
    for j, col in enumerate(df.columns, start=start_col):
        cell = ws.cell(row=start_row, column=j, value=col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, value in enumerate(row, start=start_col):
            ws.cell(row=i, column=j, value=None if pd.isna(value) else value)
    last_row = start_row + len(df)
    ref = (f"{get_column_letter(start_col)}{start_row}:"
           f"{get_column_letter(start_col + len(df.columns) - 1)}{last_row}")
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    for j, col in enumerate(df.columns, start=start_col):
        ws.column_dimensions[get_column_letter(j)].width = max(12, len(str(col)) + 4)
    return start_row, last_row


def kpi(ws, col, label, formula, number_format):
    letter = get_column_letter(col)
    ws.merge_cells(f"{letter}3:{get_column_letter(col + 1)}3")
    ws.merge_cells(f"{letter}4:{get_column_letter(col + 1)}4")
    label_cell, value_cell = ws[f"{letter}3"], ws[f"{letter}4"]
    label_cell.value, value_cell.value = label, formula
    label_cell.font, value_cell.font = KPI_LABEL_FONT, KPI_VALUE_FONT
    value_cell.number_format = number_format
    for cell in (label_cell, value_cell):
        cell.fill = KPI_FILL
        cell.alignment = Alignment(horizontal="center")


def main():
    fva = pd.read_csv(PROCESSED_DIR / "forecast_vs_actual.csv", parse_dates=["ds"])
    comparison = pd.read_csv(PROCESSED_DIR / "model_comparison.csv")
    monthly = pd.read_csv(PROCESSED_DIR / "monthly_category_sales.csv", parse_dates=["sales_month"])
    next_90 = pd.read_csv(PROCESSED_DIR / "next_90_day_forecast.csv")

    # Chart inputs.
    monthly_total = (monthly.groupby("sales_month", as_index=False)["revenue"].sum()
                     .rename(columns={"sales_month": "Month", "revenue": "Total Revenue"}))
    mape_wide = (comparison.pivot(index="product_category", columns="model", values="mape")
                 .reset_index().rename(columns={"product_category": "Category"}))
    mape_wide = mape_wide[["Category", "Prophet", "LightGBM", "Seasonal Naive"]]
    next_sorted = (next_90.sort_values("forecast_units_90d", ascending=False)
                   [["product_category", "forecast_units_90d"]]
                   .rename(columns={"product_category": "Category",
                                    "forecast_units_90d": "Forecast Units (Next 90 Days)"}))
    backtest_window = fva[fva["series"] == "Backtest"]["ds"]
    in_window = fva[fva["ds"].between(backtest_window.min(), backtest_window.max())]
    backtest_total = (in_window[in_window["series"].isin(["Actual", "Backtest"])]
                      .pivot_table(index="ds", columns="series", values="units", aggfunc="sum")
                      .reset_index().rename(columns={"ds": "Date", "Actual": "Actual Units",
                                                     "Backtest": "Backtest Units"}))

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    summary = wb.create_sheet("Summary")
    data_sheets = {
        "ForecastVsActual": fva, "ModelComparison": comparison,
        "MonthlySales": monthly, "Next90Days": next_90,
    }
    for sheet_name, df in data_sheets.items():
        write_table(wb.create_sheet(sheet_name), df, sheet_name)

    # Summary sheet: four side-by-side tables that the charts reference.
    blocks = {}
    col = 1
    for name, df in (("MonthlyTotal", monthly_total), ("MapeByModel", mape_wide),
                     ("Next90Sorted", next_sorted), ("BacktestTotal", backtest_total)):
        first, last = write_table(summary, df, name, start_col=col)
        blocks[name] = (col, first, last, len(df.columns))
        col += len(df.columns) + 1
    for r in range(2, len(monthly_total) + 2):
        summary.cell(row=r, column=1).number_format = "mmm yyyy"
    bt_col = blocks["BacktestTotal"][0]
    for r in range(2, len(backtest_total) + 2):
        summary.cell(row=r, column=bt_col).number_format = "dd mmm yyyy"

    # Dashboard: title, KPI cards (formulas), charts.
    dash["A1"] = "Sales Demand Forecasting Dashboard"
    dash["A1"].font = Font(size=20, bold=True, color="1F4E78")
    dash["A2"] = ("Prophet vs LightGBM vs a seasonal-naive baseline, each forecasting the last 90 days blind "
                  "(trained only on data before 2025-10-03). Backtest chart shows Prophet. "
                  "Source: warehouse.mart_category_daily_sales via the Airflow pipeline.")
    dash["A2"].font = Font(italic=True, color="44546A")
    kpi(dash, 1, "Total Revenue (2023-2025)", "=SUM(MonthlySales[revenue])", '"$"#,##0,,"M"')
    kpi(dash, 3, "Total Units Sold", "=SUM(MonthlySales[units_sold])", "#,##0")
    kpi(dash, 5, "Prophet Mean MAPE", '=AVERAGEIFS(ModelComparison[mape],ModelComparison[model],"Prophet")/100', "0.0%")
    kpi(dash, 7, "LightGBM Mean MAPE", '=AVERAGEIFS(ModelComparison[mape],ModelComparison[model],"LightGBM")/100', "0.0%")
    kpi(dash, 9, "Seasonal Naive Mean MAPE", '=AVERAGEIFS(ModelComparison[mape],ModelComparison[model],"Seasonal Naive")/100', "0.0%")
    kpi(dash, 11, "Forecast Units, Next 90 Days (Prophet)", "=SUM(Next90Days[forecast_units_90d])", "#,##0")
    for c in range(1, 13):
        dash.column_dimensions[get_column_letter(c)].width = 15

    def ref(block, data_cols, with_header=True):
        c0, first, last, _ = blocks[block]
        return Reference(summary, min_col=c0 + data_cols[0], max_col=c0 + data_cols[1],
                         min_row=first if with_header else first + 1, max_row=last)

    def cats(block):
        c0, first, last, _ = blocks[block]
        return Reference(summary, min_col=c0, min_row=first + 1, max_row=last)

    revenue = LineChart()
    revenue.title = "Monthly Revenue, All Categories"
    revenue.y_axis.number_format = '"$"#,##0.0,,"M"'
    revenue.add_data(ref("MonthlyTotal", (1, 1)), titles_from_data=True)
    revenue.set_categories(cats("MonthlyTotal"))
    revenue.x_axis.number_format = "mmm yy"

    mape_chart = BarChart()
    mape_chart.type, mape_chart.grouping = "bar", "clustered"
    mape_chart.title = "Holdout MAPE % by Category (lower is better)"
    mape_chart.add_data(ref("MapeByModel", (1, 3)), titles_from_data=True)
    mape_chart.set_categories(cats("MapeByModel"))

    forecast_chart = BarChart()
    forecast_chart.type = "bar"
    forecast_chart.title = "Forecast Units, Next 90 Days"
    forecast_chart.add_data(ref("Next90Sorted", (1, 1)), titles_from_data=True)
    forecast_chart.set_categories(cats("Next90Sorted"))
    forecast_chart.legend = None
    # On horizontal bar charts x_axis is the category axis; maxMin lists rows top-down,
    # and crosses='max' keeps the value axis at the bottom instead of under the title.
    for bar in (forecast_chart, mape_chart):
        bar.x_axis.scaling.orientation = "maxMin"
        bar.y_axis.crosses = "max"
    forecast_chart.y_axis.number_format = "#,##0"

    backtest = LineChart()
    backtest.title = "Backtest: Actual vs Prophet Predicted Units (All Categories, Last 90 Days)"
    backtest.add_data(ref("BacktestTotal", (1, 2)), titles_from_data=True)
    backtest.set_categories(cats("BacktestTotal"))
    backtest.x_axis.number_format = "dd mmm"

    revenue.legend = None
    for chart, anchor in ((backtest, "A6"), (mape_chart, "F6"), (revenue, "A24"), (forecast_chart, "F24")):
        # openpyxl leaves these unset, and current Excel then hides both axes and
        # colours every point differently (a 36-entry legend on a single line).
        chart.varyColors = False
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        for series in chart.series:
            series.smooth = False
        chart.width, chart.height = 17, 8.5
        dash.add_chart(chart, anchor)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}: sheets={wb.sheetnames}, charts on Dashboard={len(dash._charts)}")


if __name__ == "__main__":
    main()
